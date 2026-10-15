"""Read-only source audit and Fig. 1 deliverable validation; writes this figure's report only."""
import csv, hashlib, json, os, re, shutil, sys, zipfile
from pathlib import Path
from datetime import datetime, timezone
import xml.etree.ElementTree as ET
import numpy as np
from PIL import Image, ImageChops
from pypdf import PdfReader
from pptx import Presentation

ROOT=Path(os.path.abspath(__file__)).parents[4]
OUT=ROOT/'figures/main/Fig01'; SD=OUT/'source_data'
Image.MAX_IMAGE_PIXELS=150_000_000
def sha(p):
    h=hashlib.sha256()
    with open(p,'rb') as f:
        for b in iter(lambda:f.read(8*1024*1024),b''):h.update(b)
    return h.hexdigest()
def readj(p):return json.loads(p.read_text(encoding='utf-8-sig'))
def rows(p):
    with p.open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))
def close(a,b):return bool(np.isclose(float(a),float(b),rtol=1e-9,atol=1e-10))
def main():
    checks={}; detail={}
    provenance=readj(SD/'derivation_provenance.json')
    release=readj(ROOT/'configs/core_release_s17.json')
    checks['release_identity']=sha(ROOT/'configs/core_release_s17.json')==provenance['release_sha256']
    mismatches=[p for p,h in release['sha256'].items() if sha(ROOT/p)!=h]
    checks['all_frozen_release_files_unchanged']=not mismatches
    detail['frozen_release_files_rehashed']=len(release['sha256']);detail['frozen_mismatches']=mismatches
    checks['figure_input_hashes']=all(sha(ROOT/p)==h for p,h in provenance['source_hashes'].items())
    proof=readj(ROOT/'reports/completed_batch_verification.json')
    value_record=next(r for r in proof['prediction_files'] if Path(r['path']).name=='value_gate.npy')
    checks['value_gate_thumbnail_source_hash']=sha(ROOT/value_record['path'])==value_record['sha256']
    detail['value_gate_source']=value_record
    checks['derivation_script_preserved']=sha(OUT/'scripts/derive_source_data.py')==provenance['script_sha256']
    checks['case_rule_preserved']=sha(SD/'case_selection_rule.json')==provenance['case_rule_sha256']
    maps=rows(SD/'panel_d_h_map_metrics.csv'); quads=rows(SD/'panel_e_quadrant_shares.csv')
    display=rows(SD/'panel_f_ecdf_display.csv');cdf_errors={}
    for q in quads:
        region=q['region']; name='pooled' if region.startswith('Pooled') else region
        z=np.load(SD/f'panel_e_{name}.npz')
        checks[f'{region}_histogram_mass']=close(z['area_m2'].sum(),q['valid_area_m2']) and close(z['flood_area_m2'].sum(),q['reference_flood_area_m2'])
        checks[f'{region}_quadrant_denominators']=close(float(q['lower_entropy_higher_brier_all_m2'])/float(q['valid_area_m2']),q['quadrant_share_all_valid']) and close(float(q['lower_entropy_higher_brier_flood_m2'])/float(q['reference_flood_area_m2']),q['quadrant_share_reference_flood'])
        if name=='pooled':continue
        for model in ['sar','fusion']:
            x=np.load(SD/f'panel_f_exact_{region}_{model}.npz'); p,w,c=x['probability'],x['area_m2'],x['cdf']
            ref=next(r for r in maps if r['region']==region and r['model']==model)
            checks[f'{region}_{model}_ecdf_exact']=bool(np.all(np.diff(p)>0) and np.all(w>0) and np.all(np.diff(c)>=0) and close(c[-1],1) and close(w[p<=.1].sum()/w.sum(),ref['confident_omission_01']))
            d=[r for r in display if r['region']==region and r['model']==model]
            dp=np.array([float(r['probability']) for r in d]);dc=np.array([float(r['ecdf']) for r in d]);idx=np.searchsorted(dp,p,side='right')-1
            approx=np.where(idx>=0,dc[np.maximum(idx,0)],0);err=float(np.max(np.abs(c-approx)))
            cdf_errors[f'{region}_{model}']=err
            checks[f'{region}_{model}_display_error']=err<=.00010001
    detail['maximum_ecdf_vertical_errors']=cdf_errors
    candidates=rows(SD/'panel_c_all_eligible_cases.csv');cases=readj(SD/'panel_c_selected_cases.json')
    for case in cases:
        group=[r for r in candidates if r['region']==case['region']]
        target=float(np.quantile([float(r['largest_connected_omission_m2']) for r in group],.75))
        distance=abs(case['largest_connected_omission_m2']-target)
        checks[f"{case['region']}_deterministic_case"]=close(target,case['target_75th_percentile_m2']) and close(distance,min(abs(float(r['largest_connected_omission_m2'])-target) for r in group)) and len(group)==case['eligible_cases'] and sha(ROOT/case['label_path'])==case['label_sha256']
    export_details={}
    for suffix in ['', '_Arial']:
        stem='Fig01_main'+suffix; pdf=PdfReader(OUT/(stem+'.pdf'));page=pdf.pages[0]
        checks[stem+'_pdf']=len(pdf.pages)==1 and abs(float(page.mediabox.width)*25.4/72-180)<.01 and abs(float(page.mediabox.height)*25.4/72-220)<.01 and len(page.extract_text())>1500
        fonts=[str(v.get_object().get('/BaseFont')) for v in page['/Resources']['/Font'].values()]
        svg=ET.parse(OUT/(stem+'.svg'));ns={'s':'http://www.w3.org/2000/svg'};texts=svg.findall('.//s:text',ns)
        labels=[''.join(t.itertext()) for t in texts]
        sizes=[float(m.group(1)) for t in texts for m in [re.search(r'([\d.]+)px',t.attrib.get('style',''))] if m]
        checks[stem+'_svg_text']=len(texts)>150 and all(k in labels for k in 'abcdefgh')
        checks[stem+'_font_minimum']=bool(sizes) and min(sizes)>=7
        with Image.open(OUT/(stem+'.png')) as png,Image.open(OUT/(stem+'.tiff')) as tif:
            dimensions=png.size;dpi=tuple(map(float,png.info['dpi']));tdpi=tuple(map(float,tif.info['dpi']))
            checks[stem+'_raster_dimensions']=png.size==tif.size and all(abs(a-b)<1.1 for a,b in zip(dimensions,[180/25.4*1200,220/25.4*1200]))
            checks[stem+'_1200dpi']=all(abs(v-1200)<.02 for v in (*dpi,*tdpi))
            equal=True
            for y in range(0,png.height,512):
                box=(0,y,png.width,min(y+512,png.height))
                if ImageChops.difference(png.crop(box).convert('RGB'),tif.crop(box).convert('RGB')).getbbox():equal=False;break
            checks[stem+'_lossless_equivalence']=equal
            export_details[stem]={'pixel_dimensions':dimensions,'png_dpi':dpi,'tiff_dpi':tdpi,'pdf_fonts':fonts,'svg_text_objects':len(texts),'parsed_font_sizes':sorted(set(sizes))}
    detail['exports']=export_details
    pres=Presentation(OUT/'Fig01_editable.pptx');native=sum(s.has_text_frame for s in pres.slides[0].shapes)
    with zipfile.ZipFile(OUT/'Fig01_editable.pptx') as z:
        checks['pptx_zip_integrity']=z.testzip() is None
        svgs=[n for n in z.namelist() if n.endswith('.svg')];detail['pptx_svg_panels']=len(svgs)
    checks['pptx_editability']=len(pres.slides)==1 and native>=200 and len(svgs)==7
    detail['pptx_native_text_objects']=native;detail['pptx_shapes']=len(pres.slides[0].shapes)
    checks['pptx_actual_render']=(OUT/'assets/pptx_render/Fig01_editable.png').stat().st_size>100000 and len(PdfReader(OUT/'assets/pptx_render/Fig01_editable.pdf').pages)==1
    caption=(OUT/'Fig01_caption_en.md').read_text(encoding='utf-8-sig');detail['caption_words']=len(caption.split());checks['caption_max_350_words']=detail['caption_words']<=350
    checks['eight_panel_layout']=set(readj(OUT/'assets/panel_layout_mm.json')['panels'])==set('abcdefgh')
    total=0
    for directory,dirs,files in os.walk(ROOT):
        for name in files:total+=(Path(directory)/name).stat().st_size
    detail['project_bytes_at_validation']=total;detail['disk_free_bytes']=shutil.disk_usage(ROOT).free;checks['project_under_100GB']=total<100_000_000_000
    bound=[]
    for p in OUT.rglob('*'):
        if not p.is_file() or 'lo_profile' in p.parts or '__pycache__' in p.parts or p.name=='Fig01_validation.json' or p.suffix=='.zip':continue
        if p.parent==OUT or 'source_data' in p.parts or 'scripts' in p.parts or p.name.startswith('panel_') or p.name in ['workflow_native_commands.json','pptx_native_text_manifest.json']:
            bound.append({'path':str(p.relative_to(OUT)),'bytes':p.stat().st_size,'sha256':sha(p)})
    result={'passed':all(checks.values()),'checked_at_utc':datetime.now(timezone.utc).isoformat(),'release_sha256':provenance['release_sha256'],'checks':checks,'details':detail,'sha256_manifest':bound,
      'visual_qa':{'rounds':['Scientific evidence and original rendering','Layout, spatial zooms, legends and independent reviewer','Final-size TNR/Arial and actual PPT PDF rendering'],'resolved':['Panel e colourbar cropping','Text sizes and crowded labels','PPT theme shadows and multiline text placement','ROI connector export coordinates','Dashed supervision arrow double drawing'],'scope':'Actual rendered previews inspected by primary agent and independent visual reviewer; statistical checks are separate.'},
      'reproduction_scope':'Source-data render is self-contained in project; full raw-prediction derivation requires original frozen project files identified by hashes. Original data and evaluation remain unchanged.'}
    (OUT/'Fig01_validation.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
    print(json.dumps({'passed':result['passed'],'failed':[k for k,v in checks.items() if not v],'details':detail},indent=2))
    return 0 if result['passed'] else 1
if __name__=='__main__':sys.exit(main())
